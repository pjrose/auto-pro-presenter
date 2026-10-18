"""Stage or apply a dated ProPresenter service from a reviewed JSON manifest."""
from __future__ import annotations

import argparse
import copy
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from uuid import uuid4

from network_check import require_network

sys.path.insert(0, str(Path(__file__).parent / "vendor"))
from presentation_pb2 import Presentation
from propresenter_pb2 import PlaylistDocument
from template_pb2 import Template
from pptx import Presentation as PowerPoint
from striprtf.striprtf import rtf_to_text


def uid() -> str:
    return str(uuid4())


def rtf_escape(value: str) -> bytes:
    pieces = []
    for char in value:
        code = ord(char)
        if char in "{}\\":
            pieces.append("\\" + char)
        elif char == "\n":
            pieces.append("\\par ")
        elif 32 <= code <= 126:
            pieces.append(char)
        else:
            if code > 32767:
                code -= 65536
            pieces.append(f"\\u{code} ?")
    return "".join(pieces).encode("ascii")


def styled_rtf(text: str, font_size: int = 75, bold: bool = True) -> bytes:
    # Text attributes are also set on the ProPresenter element; this RTF keeps
    # the editable content and Unicode characters visible to the editor.
    header = ("{\\rtf1\\ansi\\ansicpg1252"
              "{\\fonttbl{\\f0\\fnil HelveticaNeue-Bold;}}"
              "{\\colortbl;\\red255\\green255\\blue255;}"
              f"\\uc1\\pard\\qc\\f0{'\\b' if bold else ''}\\fs{font_size * 2}\\cf1 ")
    return header.encode("ascii") + rtf_escape(text) + b"}"


def set_element_text(element, text: str, size: int, bold: bool = True) -> None:
    element.element.text.rtf_data = styled_rtf(text, size, bold)
    attrs = element.element.text.attributes
    attrs.font.name = "HelveticaNeue-Bold"
    attrs.font.family = "HelveticaNeue-Bold"
    attrs.font.size = size
    attrs.font.bold = bold
    attrs.text_solid_fill.red = 1
    attrs.text_solid_fill.green = 1
    attrs.text_solid_fill.blue = 1
    attrs.text_solid_fill.alpha = 1
    element.element.text.scale_behavior = 2  # scale down to fit


def set_template_text(element, text: str) -> None:
    """Replace placeholder text while retaining a theme element's RTF style."""
    data = element.element.text.rtf_data
    # Use the FIRST run: later runs contain placeholder lyrics and their own
    # paragraph formatting. Retain font/color tables and text stroke controls.
    marker = re.search(rb"\\cb\d+ ", data)
    if marker is None or not data.endswith(b"}"):
        raise ValueError(f"Unrecognized themed RTF in {element.element.name}")
    attrs = element.element.text.attributes
    capitalization = attrs.DESCRIPTOR.fields_by_name["capitalization"].enum_type
    if capitalization.values_by_number[attrs.capitalization].name == "CAPITALIZATION_ALL_CAPS":
        text = text.upper()
    # Placeholder offsets no longer address this text. They otherwise cause
    # only fragments of the new lyrics to inherit the theme's capitalization.
    attrs.ClearField("custom_attributes")
    element.element.text.rtf_data = data[:marker.end()] + rtf_escape(text) + b"}"
    element.element.text.scale_behavior = 2


def normalize_slide_size(slide) -> None:
    """Fit the entire canvas to Full HD, including an older first cue."""
    width, height = slide.size.width, slide.size.height
    if width <= 0 or height <= 0:
        raise ValueError("Cannot resize a slide with missing or invalid dimensions")
    sx, sy = 1920 / width, 1080 / height
    for item in slide.elements:
        bounds = item.element.bounds
        bounds.origin.x *= sx
        bounds.origin.y *= sy
        bounds.size.width *= sx
        bounds.size.height *= sy
    slide.size.width, slide.size.height = 1920, 1080


def format_song(doc: Presentation, lyric_template) -> None:
    """Apply the lyric theme without replacing cue IDs, arrangements or media."""
    template = copy.deepcopy(lyric_template)
    normalize_slide_size(template)
    themed = [e for e in template.elements if e.element.HasField("text")]
    if len(themed) != 1:
        raise ValueError("The selected lyric theme needs one text element")
    for cue in doc.cues:
        for action in cue.actions:
            if not action.HasField("slide") or not action.slide.HasField("presentation"):
                continue
            slide = action.slide.presentation.base_slide
            normalize_slide_size(slide)
            elements = [e for e in slide.elements if e.element.HasField("text")]
            if len(elements) > 1:
                raise ValueError(f"Review multi-text lyric slide before theming: {doc.name}, cue {cue.uuid.string}")
            for element in elements:
                text = rtf_to_text(element.element.text.rtf_data.decode("utf-8"))
                original_id = element.element.uuid.string
                element.element.CopyFrom(themed[0].element)
                element.element.uuid.string = original_id
                set_template_text(element, text)


def fresh_slide(template_slide):
    slide = copy.deepcopy(template_slide)
    normalize_slide_size(slide)
    slide.uuid.string = uid()
    for element in slide.elements:
        element.element.uuid.string = uid()
    return slide


def make_cue(source_cue, source_action, slide, name: str = ""):
    cue = copy.deepcopy(source_cue)
    cue.uuid.string = uid()
    cue.name = name
    cue.ClearField("actions")
    action = cue.actions.add()
    action.CopyFrom(source_action)
    action.uuid.string = uid()
    action.ClearField("label")
    action.slide.presentation.base_slide.CopyFrom(slide)
    return cue


def pptx_slide_texts(path: Path) -> list[str]:
    deck = PowerPoint(path)
    slides = []
    for slide in deck.slides:
        contents = [shape.text.strip() for shape in slide.shapes
                    if shape.has_text_frame and shape.text.strip()]
        slides.append("\n".join(contents).replace("\x0b", "\n").strip())
    return slides


def make_lyrics(song: dict, show: Path, prototype: Presentation,
                source_cue, source_action, lyric_template) -> Presentation:
    slides = pptx_slide_texts(Path(song["source_pptx"]))
    doc = Presentation()
    doc.application_info.CopyFrom(prototype.application_info)
    doc.uuid.string = uid()
    doc.name = song["name"]
    doc.music_key = song.get("key", "")
    doc.last_modified_date.seconds = int(datetime.now(timezone.utc).timestamp())

    background_slide = fresh_slide(lyric_template)
    for element in background_slide.elements:
        if element.element.HasField("text"):
            set_template_text(element, "")
    background = make_cue(source_cue, source_action, background_slide, "Background")
    doc.cues.add().CopyFrom(background)
    bg_group = doc.cue_groups.add()
    bg_group.group.uuid.string = uid()
    bg_group.group.name = "Background"
    bg_group.cue_identifiers.add().string = background.uuid.string

    corrections = {int(k): v for k, v in song.get("corrections", {}).items()}
    used = set()
    section_groups: dict[tuple[str, tuple[str, ...]], str] = {}
    normal_order = [bg_group.group.uuid.string]
    for section in song["sections"]:
        section_lines = []
        for number in section["slides"]:
            if number < 1 or number > len(slides) or number in used:
                raise ValueError(f"Invalid or repeated PPTX slide {number} in {song['name']}")
            used.add(number)
            lyric = corrections.get(number, slides[number - 1])
            if not lyric:
                raise ValueError(f"Empty lyric slide {number} in {song['name']}")
            section_lines.append(lyric)
        identity = (section["name"], tuple(section_lines))
        if identity in section_groups:
            normal_order.append(section_groups[identity])
            continue
        group = doc.cue_groups.add()
        group.group.uuid.string = uid()
        group.group.name = section["name"]
        section_groups[identity] = group.group.uuid.string
        normal_order.append(group.group.uuid.string)
        for lyric in section_lines:
            slide = fresh_slide(lyric_template)
            text_elements = [e for e in slide.elements if e.element.HasField("text")]
            if len(text_elements) != 1:
                raise ValueError("The selected lyric theme needs one text element")
            set_template_text(text_elements[0], lyric)
            cue = make_cue(source_cue, source_action, slide)
            doc.cues.add().CopyFrom(cue)
            group.cue_identifiers.add().string = cue.uuid.string

    remaining = [i + 1 for i, text in enumerate(slides) if text and i + 1 not in used]
    if remaining:
        raise ValueError(f"Unassigned lyric slides for {song['name']}: {remaining}")
    arrangement = doc.arrangements.add()
    arrangement.uuid.string = uid()
    arrangement.name = "Normal"
    for group_id in normal_order:
        arrangement.group_identifiers.add().string = group_id
    doc.selected_arrangement.string = arrangement.uuid.string
    return doc


def patch_existing(doc: Presentation, fixes: list[dict]) -> int:
    changed = 0
    for fix in fixes:
        before = rtf_escape(fix["from"])
        after = rtf_escape(fix["to"])
        matches = 0
        for cue in doc.cues:
            for action in cue.actions:
                if not action.HasField("slide") or not action.slide.HasField("presentation"):
                    continue
                for element in action.slide.presentation.base_slide.elements:
                    if element.element.HasField("text"):
                        data = element.element.text.rtf_data
                        count = data.count(before)
                        if count:
                            element.element.text.rtf_data = data.replace(before, after)
                            matches += count
        if matches != fix.get("expected_count", matches):
            raise ValueError(f"Text fix count differed for {doc.name}: {fix['from']} ({matches})")
        if matches == 0:
            raise ValueError(f"Text fix not found in {doc.name}: {fix['from']}")
        changed += matches
    if changed:
        doc.last_modified_date.seconds = int(datetime.now(timezone.utc).timestamp())
    return changed


def set_background_video(doc: Presentation, media_action, video_path: Path,
                         show: Path, duration: float, frame_rate: float) -> None:
    """Attach a looping Full HD video to the first (blank) background cue."""
    if not doc.cues:
        raise ValueError(f"No cues in {doc.name}")
    cue = doc.cues[0]
    if not any(a.HasField("slide") for a in cue.actions):
        raise ValueError(f"First cue has no slide in {doc.name}")
    for index in reversed(range(len(cue.actions))):
        if cue.actions[index].HasField("media"):
            del cue.actions[index]
    action = cue.actions.add()
    action.CopyFrom(media_action)
    action.uuid.string = uid()
    action.name = video_path.stem
    action.ClearField("label")
    media = action.media.element
    media.uuid.string = uid()
    media.metadata.ClearField("manufacture_name")
    media.metadata.format = "h264"
    for url in (media.url, media.video.file.local_url):
        url.absolute_string = str(video_path)
        url.local.path = video_path.relative_to(show).as_posix()
    media.video.file.ClearField("remote_properties")
    transport = media.video.transport
    transport.out_point = duration
    transport.end_point = duration
    media.video.video.frame_rate = frame_rate
    media.video.drawing.natural_size.width = 1920
    media.video.drawing.natural_size.height = 1080
    doc.last_modified_date.seconds = int(datetime.now(timezone.utc).timestamp())


def simple_slide(theme_slide, text_by_name: dict[str, str], size_by_name: dict[str, int]):
    slide = fresh_slide(theme_slide.base_slide)
    slide.draws_background_color = True
    slide.background_color.red = 0
    slide.background_color.green = 0
    slide.background_color.blue = 0
    slide.background_color.alpha = 1
    keep = [e for e in slide.elements if e.element.HasField("text") and e.element.name in text_by_name]
    slide.ClearField("elements")
    for element in keep:
        target = slide.elements.add()
        target.CopyFrom(element)
        set_element_text(target, text_by_name[element.element.name], size_by_name[element.element.name])
    return slide


def place_text(slide, name: str, x: int, y: int, width: int, height: int) -> None:
    element = next(e for e in slide.elements if e.element.name == name)
    bounds = element.element.bounds
    bounds.origin.x = x
    bounds.origin.y = y
    bounds.size.width = width
    bounds.size.height = height
    element.element.text.vertical_alignment = 1  # middle


def make_sermon(manifest: dict, prototype: Presentation, source_cue, source_action,
                black_theme: Template.Document, bible_theme_slide) -> Presentation:
    themes = {slide.name: slide for slide in black_theme.slides}
    doc = Presentation()
    doc.application_info.CopyFrom(prototype.application_info)
    doc.uuid.string = uid()
    doc.name = manifest["sermon_name"]
    doc.last_modified_date.seconds = int(datetime.now(timezone.utc).timestamp())

    title = "\n".join(manifest["sermon_title"])
    title_slide = simple_slide(themes["Title"], {"Title": title}, {"Title": 67})
    place_text(title_slide, "Title", 150, 190, 1620, 700)
    doc.cues.add().CopyFrom(make_cue(source_cue, source_action, title_slide))
    for entry in manifest["sermon_slides"]:
        if entry["type"] == "point":
            point = entry["text"]
            size = 57 if len(point) < 65 else 49
            slide = simple_slide(themes["Title"], {"Title": point}, {"Title": size})
            place_text(slide, "Title", 150, 190, 1620, 700)
        elif entry["type"] == "verse":
            ref = entry["reference"]
            verse = entry["text"]
            slide = fresh_slide(bible_theme_slide.base_slide)
            slide.draws_background_color = True
            slide.background_color.red = 0
            slide.background_color.green = 0
            slide.background_color.blue = 0
            slide.background_color.alpha = 1
            text_elements = {e.element.name: e for e in slide.elements if e.element.HasField("text")}
            if set(text_elements) != {"Verse", "Reference"}:
                raise ValueError("Central Bible Format needs Verse and Reference elements")
            set_template_text(text_elements["Verse"], verse)
            set_template_text(text_elements["Reference"], ref + " (RVR1960)")
        else:
            raise ValueError(f"Unknown sermon slide type: {entry['type']}")
        doc.cues.add().CopyFrom(make_cue(source_cue, source_action, slide))
    # ProPresenter builds its presentation viewer from cue groups. Ungrouped
    # cues parse successfully but leave the presentation area empty in the UI.
    arrangement = doc.arrangements.add()
    arrangement.uuid.string = uid()
    arrangement.name = "Normal"
    for index, cue in enumerate(doc.cues):
        group = doc.cue_groups.add()
        group.group.uuid.string = uid()
        if index == 0:
            group.group.name = "Title"
        else:
            kind = manifest["sermon_slides"][index - 1]["type"]
            group.group.name = "Point" if kind == "point" else "Scripture"
        group.cue_identifiers.add().string = cue.uuid.string
        arrangement.group_identifiers.add().string = group.group.uuid.string
    doc.selected_arrangement.string = arrangement.uuid.string
    if len(doc.cue_groups) != len(doc.cues):
        raise ValueError("Every sermon cue must belong to a cue group")
    return doc


def playlist_item(node, name: str, path: Path, show: Path, presentation: Presentation):
    item = node.items.items.add()
    item.uuid.string = uid()
    item.name = name
    url = item.presentation.document_path
    url.absolute_string = str(path)
    url.platform = 2  # Windows
    url.local.root = 10  # ProPresenter show root
    url.local.path = path.relative_to(show).as_posix()
    if presentation.arrangements:
        arrangement_id = presentation.selected_arrangement.string or presentation.arrangements[0].uuid.string
        item.presentation.arrangement.string = arrangement_id
        for arrangement in presentation.arrangements:
            if arrangement.uuid.string == arrangement_id:
                item.presentation.arrangement_name = arrangement.name
                break


def make_playlist(show: Path, date: str, songs: list[tuple[str, Path, Presentation]],
                  sermon: tuple[str, Path, Presentation] | None) -> PlaylistDocument:
    playlist_path = show / "Playlists/Library"
    if playlist_path.exists():
        source = PlaylistDocument.FromString(playlist_path.read_bytes())
    else:
        source = PlaylistDocument.FromString((Path(__file__).parent / "templates/empty-playlist").read_bytes())
    if any(p.name == date for p in source.root_node.playlists.playlists):
        raise ValueError(f"Playlist {date} already exists")
    node = source.root_node.playlists.playlists.add()
    node.uuid.string = uid()
    node.name = date
    node.items.SetInParent()
    for name, path, doc in [*songs, *([sermon] if sermon else [])]:
        playlist_item(node, name, path, show, doc)
    return source


def app_is_running() -> bool:
    if os.name != "nt":
        return False
    result = subprocess.run(["tasklist", "/FI", "IMAGENAME eq ProPresenter.exe", "/FO", "CSV"],
                            capture_output=True, text=True, check=True)
    return '"ProPresenter.exe"' in result.stdout


def apply_staged(stage: Path, show: Path, backup: Path, relative_paths: list[Path]) -> None:
    if app_is_running():
        raise RuntimeError("Close ProPresenter before applying staged files")
    if not show.is_dir() or not any(show.iterdir()):
        raise FileNotFoundError("ProPresenter show folder is missing or empty. Start ProPresenter once and confirm --show-root before applying.")
    backup.mkdir(parents=True, exist_ok=False)
    for rel in relative_paths:
        src = show / rel
        if src.exists():
            dst = backup / rel
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dst)
    applied = []
    try:
        for library_name in ("Lyrics", "Central en Espanol"):
            (show / "Libraries" / library_name).mkdir(parents=True, exist_ok=True)
        for rel in relative_paths:
            dst = show / rel
            dst.parent.mkdir(parents=True, exist_ok=True)
            temp = dst.with_name(dst.name + ".codex-new")
            shutil.copy2(stage / rel, temp)
            os.replace(temp, dst)
            applied.append(rel)
    except Exception:
        for rel in reversed(applied):
            saved = backup / rel
            if saved.exists():
                shutil.copy2(saved, show / rel)
            else:
                (show / rel).unlink(missing_ok=True)
        raise


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("manifest", type=Path)
    parser.add_argument("--show-root", type=Path, required=True)
    parser.add_argument("--stage", type=Path, required=True)
    parser.add_argument("--apply", action="store_true", help="Back up and apply, only when ProPresenter is closed")
    parser.add_argument("--api-url", help="ProPresenter API URL for the initial staging preflight")
    args = parser.parse_args()
    # Applying is deliberately offline: ProPresenter must already be closed.
    # Normal staging must fail before reading source files if the API is down.
    if not args.apply:
        require_network(args.api_url)
    show = args.show_root.resolve()
    stage = args.stage.resolve()
    manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
    date = manifest["playlist_date"]
    existing_prototypes = list((show / "Libraries/Lyrics").glob("*.pro"))
    prototype_path = existing_prototypes[0] if existing_prototypes else Path(__file__).parent / "templates/starter.pro"
    if not prototype_path.exists():
        raise FileNotFoundError("No Lyrics presentation found. Open ProPresenter once to create a show, then add a simple slide presentation in Lyrics.")
    prototype = Presentation.FromString(prototype_path.read_bytes())
    source_cue = next(c for c in prototype.cues if any(a.HasField("slide") and a.slide.HasField("presentation") for a in c.actions))
    source_action = next(a for a in source_cue.actions if a.HasField("slide") and a.slide.HasField("presentation"))

    def load_theme(name: str) -> Template.Document:
        bundled = Path(__file__).parent / "themes" / name / "Theme"
        source = bundled if bundled.exists() else show / "Themes" / name / "Theme"
        return Template.Document.FromString(source.read_bytes())

    lyric_theme = load_theme("Lyrics Centered")
    top_third = next(s.base_slide for s in lyric_theme.slides if s.name == "Top Third - White")
    black_theme = load_theme("Black")
    plain_theme = load_theme("Plain Text")
    bible_slide = next(s for s in plain_theme.slides if s.name == "Central Bible Format")
    media_action = next((a for c in prototype.cues for a in c.actions if a.HasField("media")), None)

    staged: dict[Path, bytes] = {}
    song_entries = []
    for song in manifest["songs"]:
        name = song["name"]
        path = show / "Libraries" / "Lyrics" / f"{name}.pro"
        if path.exists():
            doc = Presentation.FromString(path.read_bytes())
            if song.get("source_pptx") and not song.get("use_existing", True):
                raise ValueError(f"Presentation {name} already exists")
            if song.get("fixes"):
                patch_existing(doc, song["fixes"])
                staged[path.relative_to(show)] = doc.SerializeToString()
        else:
            if not song.get("source_pptx"):
                raise FileNotFoundError(f"No existing presentation or PPTX for {name}")
            doc = make_lyrics(song, show, prototype, source_cue, source_action, top_third)
            staged[path.relative_to(show)] = doc.SerializeToString()
        if song.get("background_video"):
            if media_action is None:
                raise ValueError("This show has no video cue to use as a model. Add one video background to a Lyrics presentation in ProPresenter, then try again. Lyrics without video can be built now.")
            video = song["background_video"]
            source_video = Path(video["source"])
            if not source_video.is_file():
                raise FileNotFoundError(source_video)
            rel_video = Path("Media/Downloads/Service Backgrounds") / source_video.name
            dest_video = show / rel_video
            staged[rel_video] = source_video.read_bytes()
            set_background_video(doc, media_action, dest_video, show,
                                 float(video["duration"]), float(video["frame_rate"]))
            staged[path.relative_to(show)] = doc.SerializeToString()
        # Reused songs also need the selected theme and a consistent canvas.
        # In particular, the first old cue can determine the presentation size.
        format_song(doc, top_third)
        staged[path.relative_to(show)] = doc.SerializeToString()
        song_entries.append((name, path, doc))

    sermon = None
    sermon_entry = None
    if manifest.get("sermon_slides"):
        missing = [s.get("reference", "(unnamed)") for s in manifest["sermon_slides"]
                   if s.get("type") == "verse" and not s.get("text", "").strip()]
        if missing:
            print("Sermon skipped: Bible text missing for " + ", ".join(missing)
                  + ". Check Reina-Valera 1960 access in ProPresenter or ask for approved verse text. The lyric playlist will still be prepared.")
        else:
            sermon_path = show / "Libraries" / "Central en Espanol" / f"{manifest['sermon_name']}.pro"
            if sermon_path.exists():
                raise ValueError(f"Sermon presentation already exists: {sermon_path}")
            sermon = make_sermon(manifest, prototype, source_cue, source_action, black_theme, bible_slide)
            staged[sermon_path.relative_to(show)] = sermon.SerializeToString()
            sermon_entry = (manifest["sermon_name"], sermon_path, sermon)
    playlist = make_playlist(show, date, song_entries, sermon_entry)
    staged[Path("Playlists/Library")] = playlist.SerializeToString()
    for theme_name in ("Black", "Lyrics Centered", "Plain Text"):
        rel = Path("Themes") / theme_name / "Theme"
        if not (show / rel).exists():
            staged[rel] = (Path(__file__).parent / "themes" / theme_name / "Theme").read_bytes()

    stage.mkdir(parents=True, exist_ok=True)
    for rel, payload in staged.items():
        target = stage / rel
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(payload)
        if rel == Path("Playlists/Library"):
            parsed = PlaylistDocument.FromString(payload)
        elif rel.suffix == ".pro":
            parsed = Presentation.FromString(payload)
        else:
            parsed = None
        if parsed is not None and parsed.SerializeToString() != payload:
            raise ValueError(f"Round-trip validation failed: {rel}")
    print(f"Staged {len(staged)} files at {stage}")
    print(f"Playlist {date}: {[x[0] for x in song_entries] + ([sermon_entry[0]] if sermon_entry else [])}")
    if sermon:
        print(f"Sermon slides: {len(sermon.cues)}")
    if args.apply:
        backup = stage / ("backup-" + datetime.now().strftime("%Y%m%d-%H%M%S"))
        apply_staged(stage, show, backup, list(staged))
        print(f"Applied to {show}; backup: {backup}")


if __name__ == "__main__":
    try:
        main()
    except (ValueError, FileNotFoundError, RuntimeError) as exc:
        print(f"Could not prepare service: {exc}", file=sys.stderr)
        sys.exit(2)
